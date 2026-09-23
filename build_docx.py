#!/usr/bin/env python3
"""Build one printable .docx per book from the chapter Markdown.

    pip install python-docx playwright
    python3 build_docx.py            # every book that has chapters
    python3 build_docx.py matthew    # one book by slug
"""
import json
import os
import re
import sys

from docx import Document
from docx.enum.section import WD_ORIENT, WD_SECTION
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt, Inches

from build_site import parse_chapter, split_notes, VERSE_RE

import io
import struct

ROOT = os.path.dirname(os.path.abspath(__file__))
OUT = os.path.join(ROOT, "docx")

BOLD_RE = re.compile(r"\*\*(.+?)\*\*")
ITAL_RE = re.compile(r"(?<!\*)\*([^*]+?)\*(?!\*)")


def add_rich(par, text):
    """Add text to a paragraph, honouring **bold** and *italic*."""
    for chunk in re.split(r"(\*\*.+?\*\*|(?<!\*)\*[^*]+?\*(?!\*))", text):
        if not chunk:
            continue
        b, i = BOLD_RE.fullmatch(chunk), ITAL_RE.fullmatch(chunk)
        run = par.add_run(b.group(1) if b else i.group(1) if i else chunk)
        run.bold, run.italic = bool(b), bool(i)


def setup(doc, title):
    normal = doc.styles["Normal"]
    normal.font.name = "Georgia"
    normal.font.size = Pt(11)
    normal.paragraph_format.space_after = Pt(8)
    for section in doc.sections:
        section.left_margin = section.right_margin = Inches(1.15)
    head = doc.add_paragraph()
    head.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = head.add_run(title)
    run.bold, run.font.size = True, Pt(26)
    sub = doc.add_paragraph()
    sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    srun = sub.add_run("A Close Rendering")
    srun.italic, srun.font.size = True, Pt(12)
    doc.add_page_break()


# Text area of a portrait page (8.5 x 11 less margins) and of a landscape one,
# each less room for a heading.
PORTRAIT = (6.2, 8.4)
LANDSCAPE = (8.7, 5.6)
CHART_PX = 760   # layout width the chart is drawn at; it reflows to fit

# Where to cut a chart into pages without slicing through a line of text.
# axis "y" cuts top to bottom inside [lo, hi); axis "x" cuts across one wide
# element. Each cut backs up from the ideal spot until no text box crosses it.
CUTS_JS = """([axis, lo, hi, ideal, sel]) => {
  const root = sel ? document.querySelector(sel) : document.body, boxes = [];
  const walk = document.createTreeWalker(root, NodeFilter.SHOW_TEXT);
  while (walk.nextNode()) {
    if (!walk.currentNode.textContent.trim()) continue;
    const r = document.createRange(); r.selectNodeContents(walk.currentNode);
    for (const b of r.getClientRects())
      boxes.push(axis === "y" ? [b.top + scrollY, b.bottom + scrollY]
                              : [b.left + scrollX, b.right + scrollX]);
  }
  const cuts = [lo];
  while (hi - cuts[cuts.length - 1] > ideal) {
    const start = cuts[cuts.length - 1];
    let v = start + ideal;
    while (v > start + ideal / 2 && boxes.some(([a, b]) => a < v && b > v)) v -= 2;
    cuts.push(v);
  }
  cuts.push(hi);
  return cuts;
}"""

# Elements that scroll sideways on the website (the Table of Nations tree).
WIDE_JS = """() => [...document.querySelectorAll('.supp *')].filter(e =>
  e.scrollWidth > e.clientWidth + 5 && getComputedStyle(e).overflowX !== 'visible'
).map((e, i) => { e.dataset.wide = i; const r = e.getBoundingClientRect();
  return [i, r.top + scrollY, r.bottom + scrollY]; })"""


def chart_slices(html):
    """Draw a supplement's HTML in a headless browser, light theme, and return
    it as page-sized PNG slices, each tagged "portrait" or "landscape".
    Anything that scrolls sideways on the site is printed whole across
    landscape pages instead of being cut off at the page edge."""
    from playwright.sync_api import sync_playwright
    out = []
    with sync_playwright() as pw:
        # CHROMIUM_PATH lets a machine point at a browser it already has.
        browser = pw.chromium.launch(executable_path=os.environ.get("CHROMIUM_PATH"))
        page = browser.new_page(viewport={"width": CHART_PX, "height": 1000},
                                device_scale_factor=2, color_scheme="light")
        page.set_content(f'<body style="margin:0"><div class="supp">{html}</div></body>',
                         wait_until="load")
        wide = page.evaluate(WIDE_JS)
        total = page.evaluate("document.documentElement.scrollHeight")
        # Portrait runs are the stretches between the wide elements.
        runs, y = [], 0
        for i, top, bottom in wide:
            runs.append(("y", y, top, None))
            runs.append(("x", 0, 0, i))
            y = bottom
        runs.append(("y", y, total, None))

        for axis, lo, hi, i in runs:
            if axis == "y":
                if hi - lo < 4:
                    continue
                ideal = int(CHART_PX * PORTRAIT[1] / PORTRAIT[0])
                cuts = page.evaluate(CUTS_JS, ["y", lo, hi, ideal, None])
                for a, b in zip(cuts, cuts[1:]):
                    if b - a > 4:
                        out.append(("portrait", page.screenshot(full_page=True,
                            clip={"x": 0, "y": a, "width": CHART_PX, "height": b - a})))
                continue
            sel = f'[data-wide="{i}"]'
            el = page.locator(sel)
            el.evaluate("e => { e.style.overflow = 'visible'; e.style.width = e.scrollWidth + 'px'; }")
            w = el.evaluate("e => e.scrollWidth")
            page.set_viewport_size({"width": w + 40, "height": 1000})
            box = el.bounding_box()
            ideal = int(box["height"] * LANDSCAPE[0] / LANDSCAPE[1])
            cuts = page.evaluate(CUTS_JS, ["x", box["x"], box["x"] + w, ideal, sel])
            for a, b in zip(cuts, cuts[1:]):
                if b - a > 4:
                    out.append(("landscape", page.screenshot(full_page=True,
                        clip={"x": a, "y": box["y"], "width": b - a, "height": box["height"]})))
            # Put the page back as it was, so the portrait runs below still
            # line up with the positions measured at the start.
            el.evaluate("e => { e.style.overflow = ''; e.style.width = ''; }")
            page.set_viewport_size({"width": CHART_PX, "height": 1000})
        browser.close()
    return out


def set_orientation(doc, landscape):
    """Start a new section (and so a new page) in the given orientation."""
    sec = doc.add_section(WD_SECTION.NEW_PAGE)
    sec.orientation = WD_ORIENT.LANDSCAPE if landscape else WD_ORIENT.PORTRAIT
    short, long_ = sorted((sec.page_width, sec.page_height))
    sec.page_width, sec.page_height = (long_, short) if landscape else (short, long_)
    sec.left_margin = sec.right_margin = Inches(1.15)
    sec.top_margin = sec.bottom_margin = Inches(1)

def build(book):
    folder = os.path.join(ROOT, "books", f"{book['order']}-{book['slug']}")
    if not os.path.isdir(folder):
        return None
    files = sorted(f for f in os.listdir(folder) if f.endswith(".md"))
    if not files:
        return None

    doc = Document()
    setup(doc, book["title"])
    first = True
    supps = []
    for name in files:
        meta, body = parse_chapter(os.path.join(folder, name))
        if meta.get("kind") == "supplement":
            # A chart. Its body is HTML for the website, so the printed edition
            # carries a picture of it, one page per slice, after the chapters.
            supps.append(meta.get("title", "Chart"))
            supps.append(body)
            continue
        if not first:
            doc.add_page_break()
        first = False
        prose, notes = split_notes(body)
        for block in re.split(r"\n\s*\n", prose):
            block = block.strip()
            if not block or block == "---":
                continue
            if block.startswith("# "):
                doc.add_heading(block[2:].strip(), level=1)
            elif block.startswith("## "):
                doc.add_heading(block[3:].strip(), level=2)
            else:
                par = doc.add_paragraph()
                m = VERSE_RE.match(block)
                if m:
                    num = par.add_run(m.group(1) + " ")
                    num.bold, num.font.size = True, Pt(8)
                    add_rich(par, m.group(2).replace("  \n", "\n"))
                else:
                    add_rich(par, block)
        if notes:
            doc.add_heading("Notes", level=2)
            for item in notes:
                for i, para in enumerate(item):
                    # first paragraph is the bullet; the rest indent under it
                    par = doc.add_paragraph(
                        style="List Bullet" if i == 0 else None)
                    if i:
                        par.paragraph_format.left_indent = Inches(0.5)
                    add_rich(par, para)
                    for run in par.runs:
                        run.font.size = Pt(9.5)

    for title, html in zip(supps[::2], supps[1::2]):
        landscape = None
        for n, (orient, png) in enumerate(chart_slices(html)):
            want = orient == "landscape"
            if want != landscape:
                set_orientation(doc, want)
                landscape = want
            elif n:
                doc.add_page_break()
            if n == 0:
                doc.add_heading(title, level=1)
            # Fit the slice to the page on whichever side runs out first.
            box_w, box_h = LANDSCAPE if want else PORTRAIT
            px_w, px_h = struct.unpack(">II", png[16:24])
            doc.add_picture(io.BytesIO(png), width=Inches(min(box_w, box_h * px_w / px_h)))

    os.makedirs(OUT, exist_ok=True)
    path = os.path.join(OUT, f"{book['order']:02d}-{book['slug']}.docx")
    doc.save(path)
    return path


def main():
    books = json.load(open(os.path.join(ROOT, "manifest.json"),
                           encoding="utf-8"))["books"]
    wanted = sys.argv[1:]
    if wanted:
        books = [b for b in books if b["slug"] in wanted]
    for book in books:
        path = build(book)
        if path:
            print("wrote", os.path.relpath(path, ROOT))


if __name__ == "__main__":
    main()
